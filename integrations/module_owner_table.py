"""本地 Excel 模块负责人映射表读取"""
import os
from typing import Dict, List, Optional
from pathlib import Path

from openpyxl import load_workbook

from integrations.config_manager import get_config


class ModuleOwnerTable:
    """模块负责人映射表"""

    def __init__(self, file_path: Optional[str] = None):
        if file_path is None:
            config = get_config()
            file_path = config.feishu.module_owner_table or "data/module_owner_mapping.xlsx"

        self.file_path = file_path
        self._data: List[Dict[str, str]] = []
        self._load()

    def _load(self):
        """加载 Excel 数据"""
        if not os.path.exists(self.file_path):
            self._data = []
            return

        try:
            wb = load_workbook(self.file_path, read_only=True)
            ws = wb.active

            # 读取表头
            headers = [cell.value for cell in ws[1]]
            if not headers:
                wb.close()
                self._data = []
                return

            # 读取数据行
            self._data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:  # 跳过空行
                    continue
                record = dict(zip(headers, row))
                self._data.append(record)

            wb.close()
        except Exception as e:
            print(f"[ModuleOwnerTable] 加载 Excel 失败: {type(e).__name__}: {e}")
            self._data = []

    def get_owner(self, module_id: str) -> Optional[Dict[str, str]]:
        """
        根据模块ID获取负责人信息

        Args:
            module_id: 模块ID，如 "intent_classifier"

        Returns:
            负责人信息字典，包含 name, open_id, email 等
        """
        for record in self._data:
            if record.get("模块ID") == module_id:
                return {
                    "name": record.get("负责人姓名"),
                    "open_id": record.get("飞书Open ID"),
                    "email": record.get("邮箱"),
                    "phone": record.get("手机号"),
                    "department": record.get("部门"),
                    "level": record.get("级别"),
                    "status": record.get("状态"),
                    "remark": record.get("备注"),
                }
        return None

    def get_all_modules(self) -> List[Dict[str, str]]:
        """获取所有模块信息"""
        return self._data.copy()


if __name__ == "__main__":
    # 测试读取
    table = ModuleOwnerTable("data/module_owner_mapping.xlsx")

    print("\n所有模块:")
    for module in table.get_all_modules():
        print(f"  - {module['模块ID']}: {module['负责人姓名']}")

    print("\n查询 'intent_classifier':")
    owner = table.get_owner("intent_classifier")
    print(f"  {owner}")